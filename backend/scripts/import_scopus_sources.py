"""Build the source-coverage ISSN snapshot from Elsevier's Scopus source list.

Usage: python scripts/import_scopus_sources.py sources.xlsx \
    --output data/scopus_issns.json --compare data/scopus_issns.json

Includes the main source list and serial conference profiles, including inactive
and discontinued historical coverage. This is not proof that a particular paper
or the current year is indexed. Accepted-only titles and ISBN-only conference
proceedings are deliberately excluded. The input workbook is opened read-only.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def normalize_issn(value: Any) -> str | None:
    """Normalize Excel numbers and ISSN strings without inventing blank ISSNs.

    Checksum mismatches are reported separately, not silently removed from the
    official source. Numeric Excel cells may have lost their leading zeroes.
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if not 0 < value < 10**8 or not math.isfinite(value) or value != int(value):
            return None
        compact = str(int(value)).zfill(8)
    elif isinstance(value, str):
        compact = value.strip().upper()
        if re.fullmatch(r"[0-9]{4}-[0-9]{3}[0-9X]", compact):
            compact = compact.replace("-", "")
    else:
        return None
    if compact == "00000000" or not re.fullmatch(r"[0-9]{7}[0-9X]", compact):
        return None
    return f"{compact[:4]}-{compact[4:]}"


def valid_checksum(issn: str) -> bool:
    compact = issn.replace("-", "")
    digits = [10 if char == "X" else int(char) for char in compact]
    return sum(digit * weight for digit, weight in zip(digits, range(8, 0, -1))) % 11 == 0


def collect_issns(workbook: Any) -> tuple[set[str], dict[str, Any]]:
    """Select only source-coverage sheets; fail if their schema has changed."""
    main_sheets = [name for name in workbook.sheetnames if name.startswith("Scopus Sources ")]
    serial_sheet = "Serial Conf. Proc. with Profile"
    if len(main_sheets) != 1 or serial_sheet not in workbook.sheetnames:
        raise ValueError("Expected one 'Scopus Sources ...' sheet and the serial conference profile sheet")

    all_issns: set[str] = set()
    sheet_reports: list[dict[str, Any]] = []
    for name in [main_sheets[0], serial_sheet]:
        rows = workbook[name].iter_rows(values_only=True)
        header = next(rows)
        columns = {str(value).strip(): index for index, value in enumerate(header) if value is not None}
        required = {"Sourcerecord ID", "Source Title", "ISSN", "EISSN"}
        if not required.issubset(columns):
            raise ValueError(f"Missing required columns in {name!r}: {sorted(required - columns.keys())}")

        sheet_issns: set[str] = set()
        rejected: list[str] = []
        rejected_count = 0
        data_rows = 0
        for row in rows:
            if not any(value is not None and value != "" for value in row):
                continue
            data_rows += 1
            for column in ("ISSN", "EISSN"):
                value = row[columns[column]]
                issn = normalize_issn(value)
                if issn:
                    sheet_issns.add(issn)
                elif value is not None and str(value).strip():
                    rejected_count += 1
                    if len(rejected) < 10:
                        rejected.append(str(value))
        if not sheet_issns:
            raise ValueError(f"No valid ISSN values in {name!r}; refusing an empty import")
        all_issns.update(sheet_issns)
        sheet_reports.append({
            "sheet": name,
            "data_rows": data_rows,
            "unique_issns": len(sheet_issns),
            "rejected_nonempty_cells": rejected_count,
            "rejected_examples": rejected,
        })
    checksum_mismatches = sorted(issn for issn in all_issns if not valid_checksum(issn))
    return all_issns, {
        "sheets": sheet_reports,
        "unique_issns": len(all_issns),
        "checksum_mismatches": len(checksum_mismatches),
        "checksum_mismatch_examples": checksum_mismatches[:10],
        "coverage": "main source list and serial conferences, including historical/inactive sources",
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, help="Write sorted JSON; omitted for a read-only audit")
    parser.add_argument("--compare", type=Path, help="Report changes against an existing JSON snapshot")
    args = parser.parse_args()
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    try:
        issns, report = collect_issns(workbook)
    finally:
        workbook.close()
    report["source_file"] = args.workbook.name
    with args.workbook.open("rb") as source:
        report["source_sha256"] = hashlib.file_digest(source, "sha256").hexdigest()
    if args.compare:
        previous = set(json.loads(args.compare.read_text(encoding="utf-8")))
        report.update({
            "previous_unique_issns": len(previous),
            "added": len(issns - previous),
            "removed": len(previous - issns),
            "removed_examples": sorted(previous - issns)[:20],
        })
    if args.output:
        args.output.write_text(json.dumps(sorted(issns), ensure_ascii=False) + "\n", encoding="utf-8")
        report["output"] = str(args.output)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
